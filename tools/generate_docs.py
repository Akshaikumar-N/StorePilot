import os
from fpdf import FPDF
from pptx import Presentation
from langchain.tools import tool
from db.database import get_engine
from sqlalchemy import text

@tool
def generate_invoice_pdf(bill_id: int) -> str:
    """Generates a PDF invoice for a given bill ID and returns the local file path."""
    engine = get_engine()
    
    with engine.connect() as conn:
        bill_res = conn.execute(text("SELECT * FROM bills WHERE id = :id"), {"id": bill_id}).fetchone()
        if not bill_res:
            return f"Error: Bill {bill_id} not found."
            
        bill = bill_res._mapping
        items_res = conn.execute(text("SELECT * FROM bill_items WHERE bill_id = :id"), {"id": bill_id}).fetchall()
        items = [i._mapping for i in items_res]
    
    pdf = FPDF()
    pdf.add_page()
    

    pdf.set_font('helvetica', 'B', 16)
    pdf.cell(0, 10, 'Kirana Store - Tax Invoice', new_x="LMARGIN", new_y="NEXT", align='C')
    
    pdf.set_font('helvetica', '', 12)
    pdf.cell(0, 10, f'Bill ID: {bill["id"]}', new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 10, f'Date: {bill["created_at"]}', new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 10, f'Payment Mode: {str(bill["payment_mode"]).upper()}', new_x="LMARGIN", new_y="NEXT")
    if bill.get("customer_name"):
        pdf.cell(0, 10, f'Customer: {bill["customer_name"]}', new_x="LMARGIN", new_y="NEXT")
    
    pdf.ln(10)
    

    pdf.set_font('helvetica', 'B', 12)
    pdf.cell(60, 10, 'Item', border=1)
    pdf.cell(30, 10, 'Qty', border=1)
    pdf.cell(30, 10, 'Price', border=1)
    pdf.cell(30, 10, 'GST', border=1)
    pdf.cell(40, 10, 'Total', border=1, new_x="LMARGIN", new_y="NEXT")
    

    pdf.set_font('helvetica', '', 12)
    for item in items:
        pdf.cell(60, 10, item['product_name'][:20], border=1)
        pdf.cell(30, 10, str(item['quantity']), border=1)
        pdf.cell(30, 10, f"Rs {item['price']}", border=1)
        pdf.cell(30, 10, f"Rs {item['tax_amount']:.2f}", border=1)
        pdf.cell(40, 10, f"Rs {item['total']:.2f}", border=1, new_x="LMARGIN", new_y="NEXT")
        
    pdf.ln(10)
    pdf.set_font('helvetica', 'B', 12)
    pdf.cell(150, 10, 'Total Tax (GST):', align='R')
    pdf.cell(40, 10, f"Rs {bill['tax_amount']:.2f}", new_x="LMARGIN", new_y="NEXT")
    
    pdf.cell(150, 10, 'Grand Total:', align='R')
    pdf.cell(40, 10, f"Rs {bill['total_amount']:.2f}", new_x="LMARGIN", new_y="NEXT")
    
    os.makedirs("artifacts", exist_ok=True)
    filename = f"artifacts/invoice_{bill['id']}.pdf"
    pdf.output(filename)
    return f"PDF Invoice generated at: {filename}"

@tool
def generate_analysis_deck() -> str:
    """Generates a PPTX analysis deck of store sales and returns the file path."""
    engine = get_engine()
    
    with engine.connect() as conn:
        bills = [b._mapping for b in conn.execute(text("SELECT * FROM bills")).fetchall()]
        products = [p._mapping for p in conn.execute(text("SELECT * FROM products")).fetchall()]
    
    prs = Presentation()
    

    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Store Analysis Report"
    slide.placeholders[1].text = f"Generated from {len(bills)} bills and {len(products)} products"
    

    total_revenue = sum(b['total_amount'] for b in bills)
    total_tax = sum(b['tax_amount'] for b in bills)
    

    bullet_slide_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    slide2.shapes.title.text = "Executive Summary"
    tf = slide2.shapes.placeholders[1].text_frame
    tf.text = f"Total Revenue: Rs {total_revenue:.2f}"
    tf.add_paragraph().text = f"Total GST Collected: Rs {total_tax:.2f}"
    

    slide3 = prs.slides.add_slide(bullet_slide_layout)
    slide3.shapes.title.text = "Low Stock Alerts"
    tf3 = slide3.shapes.placeholders[1].text_frame
    
    low_stock = [p for p in products if p['stock'] <= p['reorder_level']]
    if not low_stock:
        tf3.text = "All items have sufficient stock."
    else:
        tf3.text = "Items needing reorder:"
        for p in low_stock:
            tf3.add_paragraph().text = f"- {p['name']}: {p['stock']} left (Reorder at: {p['reorder_level']})"
            
    os.makedirs("artifacts", exist_ok=True)
    filename = "artifacts/analysis_deck.pptx"
    prs.save(filename)
    return f"PPTX Analysis Deck generated at: {filename}"
